"""
Lectura de horizontes en formato xlsx (Excel).

Convierte la primera hoja del workbook a la misma lista de líneas de texto que
produce el reader de texto plano, preservando la semántica del awk legacy:

- Fila con valor en la primera celda que contiene letras → header de horizonte.
- Fila completamente vacía → ignorada.
- Fila que contiene '!' en alguna celda → ignorada (comentario legacy).
- Resto de filas → línea de datos con las celdas no-nulas unidas por un espacio.

Se usa `read_only=True` por si el archivo es grande (streaming, no carga toda
la hoja en memoria) y `data_only=True` para que las fórmulas devuelvan el
último valor cacheado en el workbook en lugar de la expresión.
"""
from __future__ import annotations

from io import BytesIO

from loguru import logger

from modules.routines.horizontes_split.errors import HorizontesSplitParseError

XLSX_MAGIC = b"PK\x03\x04"  # También usado por docx/pptx/jar: openpyxl lo rechaza si no es xlsx.


def is_xlsx_payload(raw: bytes) -> bool:
    return len(raw) >= 4 and raw[:4] == XLSX_MAGIC


def _stringify(value: object) -> str:
    """Representación estable de una celda:
    - None → cadena vacía
    - float entero (1.0) → '1' (para no ensuciar los .dat con '.0')
    - float con decimales → repr corto
    - bool → 'True' / 'False' (raro en archivos geológicos)
    - el resto → str(value).strip()
    """
    if value is None:
        return ""
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return repr(value)
    return str(value).strip()


def load_horizontes_lines_from_xlsx(raw: bytes) -> list[str]:
    """
    Retorna las líneas útiles extraídas del xlsx en el mismo formato que el
    reader de texto plano, listas para `group_lines_by_horizon(...)`.
    """
    try:
        from openpyxl import load_workbook
    except ImportError as e:  # pragma: no cover
        raise HorizontesSplitParseError(
            "openpyxl not available to read xlsx input"
        ) from e

    try:
        wb = load_workbook(
            filename=BytesIO(raw),
            read_only=True,
            data_only=True,
        )
    except Exception as e:
        raise HorizontesSplitParseError(
            f"Cannot open xlsx workbook: {e}"
        ) from e

    try:
        sheetnames = wb.sheetnames
        if not sheetnames:
            raise HorizontesSplitParseError("xlsx workbook has no sheets")
        ws = wb[sheetnames[0]]
        logger.info(
            "horizontes_split: reading xlsx sheet={!r} ({} sheets total)",
            sheetnames[0],
            len(sheetnames),
        )

        out: list[str] = []
        for row in ws.iter_rows(values_only=True):
            cells = [_stringify(v) for v in row]
            # Línea totalmente vacía → descartar (equivale a `sed '/^ *$/d'`).
            if not any(cell for cell in cells):
                continue
            # Contiene '!' en cualquier celda → descartar (equivale a `sed '/\!/d'`).
            if any("!" in cell for cell in cells):
                continue
            # Une celdas no-vacías con un espacio simple, análogo al formato `$1 $2 $3` de awk.
            line = " ".join(c for c in cells if c)
            if line.strip():
                out.append(line)
    finally:
        try:
            wb.close()
        except Exception:
            pass

    if not out:
        raise HorizontesSplitParseError(
            "xlsx sheet has no usable rows after cleanup"
        )
    return out
